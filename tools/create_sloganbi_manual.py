from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "SloganBi用户操作手册.docx"
LOGO = ROOT / "apps" / "web" / "public" / "images" / "sloganbi-logo.png"
HOME = ROOT / "artifacts" / "dashboard-home-implemented-1920x960.jpg"

INK = "12233F"
BLUE = "2476C5"
DEEP_BLUE = "174C85"
PALE_BLUE = "EAF3FC"
PALE_GRAY = "F3F6F9"
MUTED = "64748B"
WHITE = "FFFFFF"
# This is the system Simplified Chinese family used by the document renderer.
FONT = "Hiragino Sans GB"
CONTENT_DXA = 9360


def set_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color="D7E1EC", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.append(OxmlElement("w:keepNext"))


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_font(run, 9, MUTED)


def add_numbering(doc):
    numbering = doc.part.numbering_part.element
    def add_definition(abstract_id, num_id, is_bullet):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start"); start.set(qn("w:val"), "1"); lvl.append(start)
        fmt = OxmlElement("w:numFmt"); fmt.set(qn("w:val"), "bullet" if is_bullet else "decimal"); lvl.append(fmt)
        text = OxmlElement("w:lvlText"); text.set(qn("w:val"), "•" if is_bullet else "%1."); lvl.append(text)
        jc = OxmlElement("w:lvlJc"); jc.set(qn("w:val"), "left"); lvl.append(jc)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab"); tab.set(qn("w:val"), "num"); tab.set(qn("w:pos"), "540"); tabs.append(tab); p_pr.append(tabs)
        ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "540"); ind.set(qn("w:hanging"), "270"); p_pr.append(ind)
        spacing = OxmlElement("w:spacing"); spacing.set(qn("w:after"), "80"); spacing.set(qn("w:line"), "300"); spacing.set(qn("w:lineRule"), "auto"); p_pr.append(spacing)
        lvl.append(p_pr)
        abstract.append(lvl)
        numbering.append(abstract)
        n = OxmlElement("w:num"); n.set(qn("w:numId"), str(num_id))
        a = OxmlElement("w:abstractNumId"); a.set(qn("w:val"), str(abstract_id)); n.append(a); numbering.append(n)
    add_definition(99, 99, False)
    add_definition(100, 100, True)
    return 99, 100


def set_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId"); num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl); num_pr.append(num); p_pr.append(num_pr)


def add_para(doc, text="", style=None, size=None, color=None, bold=None, align=None, before=None, after=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        set_font(r, size, color, bold)
    if align is not None:
        p.alignment = align
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def add_bullet(doc, text, bullet_id):
    p = add_para(doc, text, size=10.5, after=4)
    set_num(p, bullet_id)
    return p


def add_step(doc, text, number_id):
    p = add_para(doc, text, size=10.5, after=4)
    set_num(p, number_id)
    return p


def add_note(doc, label, text, fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_DXA])
    set_table_borders(table, "B9D8F4", "8")
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + "  ")
    set_font(r, 10.5, DEEP_BLUE, True)
    r = p.add_run(text)
    set_font(r, 10.5, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_two_col_steps(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [1650, 7710])
    set_table_borders(table)
    hdr = table.rows[0].cells
    for cell, text in zip(hdr, ("步骤", "操作说明")):
        shade(cell, "E8EEF5")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); set_font(r, 10.5, INK, True)
    for i, (step, detail) in enumerate(rows, 1):
        cells = table.add_row().cells
        p = cells[0].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(i)); set_font(r, 10.5, BLUE, True)
        p = cells[1].paragraphs[0]
        r = p.add_run(step + "："); set_font(r, 10.5, INK, True)
        r = p.add_run(detail); set_font(r, 10.5, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1); section.bottom_margin = Inches(1)
    section.left_margin = Inches(1); section.right_margin = Inches(1)
    section.header_distance = Inches(0.492); section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT; normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT); normal._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (("Heading 1", 16, BLUE, 18, 10), ("Heading 2", 13, BLUE, 14, 7), ("Heading 3", 12, DEEP_BLUE, 10, 5)):
        style = doc.styles[name]
        style.font.name = FONT; style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT); style._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
        style.font.size = Pt(size); style.font.color.rgb = RGBColor.from_string(color); style.font.bold = True
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    number_id, bullet_id = add_numbering(doc)

    # Customer-pack cover
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(20); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run().add_picture(str(LOGO), width=Inches(2.35))
    p = add_para(doc, "用户操作手册", size=31, color=INK, bold=True, before=40, after=8)
    p = add_para(doc, "从数据上传、看板制作到发布分享", size=15, color=MUTED, after=28)
    add_note(doc, "适用对象", "SloganBi 业务使用者、数据分析人员及看板维护人员。", "EFF6FF")
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    meta = doc.add_table(rows=4, cols=2)
    set_table_geometry(meta, [2200, 7160]); set_table_borders(meta)
    for row, (label, value) in zip(meta.rows, (("产品名称", "SloganBi"), ("文档版本", "V1.0"), ("发布日期", "2026 年 8 月"), ("文档用途", "上线交付与日常使用参考"))):
        shade(row.cells[0], PALE_GRAY)
        a = row.cells[0].paragraphs[0].add_run(label); set_font(a, 10.5, INK, True)
        b = row.cells[1].paragraphs[0].add_run(value); set_font(b, 10.5, INK)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Contents
    add_para(doc, "阅读导航", style="Heading 1")
    for item in ("1. 快速开始", "2. 登录与账号", "3. 看板中心", "4. 数据集管理", "5. 创建与配置看板", "6. 编辑、保存与发布", "7. 常见问题", "8. 支持与反馈"):
        add_bullet(doc, item, bullet_id)
    add_note(doc, "建议", "首次使用可直接阅读“快速开始”，并在操作时对照后续章节查阅具体功能。")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    add_para(doc, "1. 快速开始", style="Heading 1")
    add_para(doc, "完成一次看板制作和发布，只需按以下流程操作。", size=10.5)
    add_two_col_steps(doc, [
        ("登录系统", "通过系统入口访问 SloganBi；首次使用请先注册账号。"),
        ("新建看板", "在看板中心点击“新建仪表板”，进入编辑器。"),
        ("上传数据", "在编辑器顶部打开“数据集”，导入 CSV 或 XLSX 文件。"),
        ("添加组件", "从左侧组件库选择图表、指标或表格，单击或拖放至画布。"),
        ("绑定字段", "选中组件，在右侧“字段”页选择数据集并完成维度、度量等字段绑定。"),
        ("保存并发布", "检查预览效果后，点击“保存并发布”，再复制发布链接分享。"),
    ])
    add_note(doc, "操作提示", "每完成一项配置，请留意编辑器顶部的保存状态；显示“已保存”后再继续下一项操作。")

    add_para(doc, "2. 登录与账号", style="Heading 1")
    add_para(doc, "2.1 注册和登录", style="Heading 2")
    add_step(doc, "在浏览器打开系统入口，进入登录页。", number_id)
    add_step(doc, "首次使用时切换至“注册”，填写账号和密码。账号长度为 4–40 个字符；密码至少 8 位，且须同时包含大写字母、小写字母、数字和符号。", number_id)
    add_step(doc, "注册成功后系统自动进入“看板中心”；已有账号可直接登录。", number_id)
    add_bullet(doc, "勾选“记住我”可保留当前登录状态。", bullet_id)
    add_para(doc, "2.2 账号设置", style="Heading 2")
    add_para(doc, "点击右上角头像进入“账号设置”，可修改展示名称、头像地址和密码，并管理已登录设备。", size=10.5)

    add_para(doc, "3. 看板中心", style="Heading 1")
    add_para(doc, "看板中心用于查看和管理本人创建的仪表板。", size=10.5)
    if HOME.exists():
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(HOME), width=Inches(6.45))
        cap = add_para(doc, "图 1  看板中心示例", size=9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    add_bullet(doc, "新建仪表板：点击右上角“新建仪表板”。系统会创建“未命名看板”并打开编辑器。", bullet_id)
    add_bullet(doc, "搜索：在“搜索看板”中输入名称关键词，快速筛选看板。", bullet_id)
    add_bullet(doc, "继续编辑：点击看板卡片，或通过卡片右上角“更多操作”选择“编辑”。", bullet_id)
    add_bullet(doc, "发布与下线：在“更多操作”中发布、打开发布页，或下线已发布看板。", bullet_id)
    add_bullet(doc, "删除：选择“删除”并确认后，将同时删除看板及其发布快照，且无法恢复。", bullet_id)
    add_note(doc, "注意", "删除不可恢复。若看板已被分享，请先通知查看方并确认不再需要访问。", "FFF7E8")

    add_para(doc, "4. 数据集管理", style="Heading 1")
    add_para(doc, "4.1 上传数据", style="Heading 2")
    add_step(doc, "在编辑器顶部点击“数据集”，打开数据集管理窗口。", number_id)
    add_step(doc, "点击“选择文件”，选择 CSV 或 XLSX（Excel）文件。", number_id)
    add_step(doc, "系统读取表头并自动推断字段类型；上传完成后可查看行数、字段和数据预览。", number_id)
    add_step(doc, "在组件右侧“字段”页的“数据集”下拉框中选择已上传的数据集。", number_id)
    add_para(doc, "4.2 数据使用注意事项", style="Heading 2")
    add_bullet(doc, "仅支持 CSV、XLSX 格式。建议第一行使用清晰且唯一的字段名称。", bullet_id)
    add_bullet(doc, "上传的数据会保存到系统中，刷新浏览器后仍可使用。", bullet_id)
    add_bullet(doc, "删除数据集可能导致已绑定的图表无法读取数据，请先确认影响范围。", bullet_id)
    add_bullet(doc, "若数据集结构发生变化，系统会提示字段版本变化；请重新检查相关图表的字段绑定。", bullet_id)

    add_para(doc, "5. 创建与配置看板", style="Heading 1")
    add_para(doc, "5.1 创建一个看板", style="Heading 2")
    add_two_col_steps(doc, [
        ("命名", "点击顶部看板名称旁的编辑图标，输入便于识别的名称。"),
        ("添加组件", "从左侧组件库单击添加，或拖动组件到画布。"),
        ("绑定数据", "选中组件后，在右侧“字段”页选择数据集和所需字段。"),
        ("调整展示", "在“显示”页设置标题、卡片和图表样式；在“分析”页设置筛选条件。"),
        ("布局", "拖动组件移动位置，拖动边缘调整大小；需要时使用顶部“一键整理”。"),
    ])
    add_para(doc, "5.2 可用组件", style="Heading 2")
    add_bullet(doc, "内容：看板信息栏、复合分析。", bullet_id)
    add_bullet(doc, "表格：明细表、热力图。", bullet_id)
    add_bullet(doc, "指标：指标看板、指标洞察、指标趋势、进度条、进度与指标、目标完成率、仪表盘。", bullet_id)
    add_bullet(doc, "图表：柱图、条形图、柱状折线组合图、堆积柱图、百分比柱图、环形柱图、排行榜、饼图、环形图、玫瑰图。", bullet_id)
    add_note(doc, "限制", "“看板信息栏”每个看板最多添加一个。")
    add_para(doc, "5.3 字段绑定与筛选", style="Heading 2")
    add_para(doc, "不同组件需要的字段不同。以柱图为例，通常需要一个分类维度和一个数值度量；饼图通常需要分类维度和数值度量。也可以从右侧数据字段列表直接拖拽字段至组件。", size=10.5)
    add_bullet(doc, "“显示”页用于调整标题、卡片和图表样式。", bullet_id)
    add_bullet(doc, "“分析”页可为当前组件设置日期筛选和查询条件。", bullet_id)
    add_bullet(doc, "看板信息栏可配置全局筛选；必须关联目标图表后才会生效。", bullet_id)
    add_bullet(doc, "复合分析可作为多个子图表的容器，并配置容器展示和相关筛选。", bullet_id)

    add_para(doc, "6. 编辑、保存与发布", style="Heading 1")
    add_para(doc, "6.1 编辑与快捷键", style="Heading 2")
    add_bullet(doc, "拖动组件可改变位置；拖动边缘可调整尺寸。", bullet_id)
    add_bullet(doc, "“一键整理”会按当前阅读顺序紧凑排列顶层组件，不会改变组件尺寸。", bullet_id)
    add_bullet(doc, "撤销/重做：Ctrl/Cmd + Z 撤销；Ctrl/Cmd + Shift + Z 重做。", bullet_id)
    add_bullet(doc, "保存：Ctrl/Cmd + S。删除选中组件：Delete 或 Backspace。", bullet_id)
    add_para(doc, "6.2 保存和冲突处理", style="Heading 2")
    add_para(doc, "编辑器会显示“正在保存”“已保存”“有未保存更改”或“保存失败”等状态。发生保存冲突时，可选择重新加载系统中的最新版本，或将本地内容复制为新看板后继续编辑。", size=10.5)
    add_para(doc, "6.3 预览、发布与分享", style="Heading 2")
    add_step(doc, "点击顶部“预览”，在不发布的情况下检查当前看板效果；预览页可返回继续编辑。", number_id)
    add_step(doc, "确认无误后点击“保存并发布”。系统先保存当前修改，再生成供查看的发布快照。", number_id)
    add_step(doc, "在预览页点击“分享”复制发布页链接；也可从看板中心“更多操作”中打开发布页。", number_id)
    add_bullet(doc, "发布页为只读模式。查看者可使用已配置的全局筛选、日期筛选等交互，但不能编辑看板。", bullet_id)
    add_bullet(doc, "修改已发布看板后，需要再次点击“保存并发布”才能更新发布快照。", bullet_id)
    add_bullet(doc, "在看板中心选择“下线发布页”后，已分享的发布链接将不再可用。", bullet_id)

    add_para(doc, "7. 常见问题", style="Heading 1")
    faqs = doc.add_table(rows=1, cols=2)
    set_table_geometry(faqs, [3000, 6360]); set_table_borders(faqs)
    for cell, title in zip(faqs.rows[0].cells, ("问题", "处理建议")):
        shade(cell, "E8EEF5")
        p = cell.paragraphs[0]; r = p.add_run(title); set_font(r, 10.5, INK, True)
    for q, a in (
        ("图表没有数据", "依次检查：是否已上传或选择数据集、字段绑定是否完整、字段类型是否正确、日期或查询筛选是否导致无匹配数据。"),
        ("发布页看不到最新修改", "请在编辑器中先点击“保存”，再点击“保存并发布”。仅保存草稿不会更新发布页。"),
        ("无法注册或修改密码", "注册和新密码均需至少 8 位，并同时包含大写字母、小写字母、数字和符号。"),
        ("数据集更新后图表异常", "检查字段版本变化提示，并重新确认受影响图表的数据集和字段绑定。"),
        ("误删了看板", "看板及发布快照删除后无法恢复。请通过日常命名、发布前复核和权限管理降低误删风险。"),
    ):
        cells = faqs.add_row().cells
        r = cells[0].paragraphs[0].add_run(q); set_font(r, 10.5, INK, True)
        r = cells[1].paragraphs[0].add_run(a); set_font(r, 10.5, INK)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_para(doc, "8. 支持与反馈", style="Heading 1")
    add_para(doc, "需要协助时，请准备以下信息后联系系统支持人员，以便更快定位问题：", size=10.5)
    add_bullet(doc, "问题发生时间、操作账号和所处页面。", bullet_id)
    add_bullet(doc, "相关看板名称、数据集名称或发布链接。", bullet_id)
    add_bullet(doc, "操作步骤、报错提示，以及必要的页面截图。", bullet_id)
    add_note(doc, "文档维护", "本手册随 SloganBi 产品功能迭代更新。请以系统当前页面和最新发布版本为准。")
    add_para(doc, "上线使用自查", style="Heading 2")
    add_bullet(doc, "已确认业务数据中不包含不应在发布页展示的敏感信息。", bullet_id)
    add_bullet(doc, "已逐个检查图表字段、筛选条件和指标口径。", bullet_id)
    add_bullet(doc, "已在预览页验证桌面端展示效果，并完成“保存并发布”。", bullet_id)
    add_bullet(doc, "已将发布链接发送给正确的查看对象，并确认其可正常访问。", bullet_id)

    # Running header/footer on all pages.
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = header.add_run("SloganBi  |  用户操作手册")
    set_font(r, 9, MUTED, True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("SloganBi  ·  ")
    set_font(r, 9, MUTED)
    add_page_number(footer)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "SloganBi 用户操作手册"
    doc.core_properties.subject = "业务使用方上线操作指南"
    doc.core_properties.author = "SloganBi"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
